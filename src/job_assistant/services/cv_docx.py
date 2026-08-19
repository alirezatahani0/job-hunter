from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_FONT = "Calibri"


def markdown_to_docx(markdown: str) -> bytes:
    """Turn resume markdown into a single-column ATS-friendly .docx (not PDF)."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    saw_title = False
    for raw in markdown.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line or line == "---":
            continue
        if line.startswith("### "):
            _add_heading(doc, line[4:], size=12, space_before=10, space_after=2)
        elif line.startswith("## "):
            _add_heading(doc, line[3:], size=13, space_before=12, space_after=4, all_caps=True)
        elif line.startswith("# "):
            _add_heading(
                doc,
                line[2:],
                size=20,
                space_before=0,
                space_after=4,
                align=WD_ALIGN_PARAGRAPH.CENTER if not saw_title else WD_ALIGN_PARAGRAPH.LEFT,
            )
            saw_title = True
        elif line.startswith(("- ", "* ")):
            _add_paragraph(doc, line[2:], bullet=True)
        else:
            is_contact = line.startswith("**") and ":" in line
            _add_paragraph(
                doc,
                line,
                size=10 if is_contact else 11,
                align=WD_ALIGN_PARAGRAPH.CENTER if is_contact else WD_ALIGN_PARAGRAPH.LEFT,
                space_after=1 if is_contact else 6,
            )

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_heading(
    doc: Document,
    text: str,
    *,
    size: int,
    space_before: int,
    space_after: int,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    all_caps: bool = False,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    _spacing(paragraph, before=space_before, after=space_after)
    label = text.strip().upper() if all_caps else text.strip()
    run = paragraph.add_run(_plain(label))
    _font(run, size=size, bold=True)


def _add_paragraph(
    doc: Document,
    text: str,
    *,
    bullet: bool = False,
    size: int = 11,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    space_after: int = 4,
) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if bullet else None)
    paragraph.alignment = align
    _spacing(paragraph, before=0, after=space_after)
    if bullet:
        paragraph.paragraph_format.left_indent = Inches(0.25)
    _add_runs(paragraph, text, size=size)


def _add_runs(paragraph, text: str, *, size: int) -> None:
    cursor = 0
    for match in _BOLD.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(_plain(text[cursor : match.start()]))
            _font(run, size=size)
        run = paragraph.add_run(_plain(match.group(1)))
        _font(run, size=size, bold=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(_plain(text[cursor:]))
        _font(run, size=size)
    elif not text:
        _font(paragraph.add_run(""), size=size)


def _plain(text: str) -> str:
    return text.replace("\\", "").strip()


def _spacing(paragraph, *, before: int, after: int) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _font(run, *, size: int, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = _FONT
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), _FONT)
    rFonts.set(qn("w:hAnsi"), _FONT)
    rFonts.set(qn("w:cs"), _FONT)
