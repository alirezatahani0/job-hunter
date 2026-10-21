from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from fpdf import FPDF

_REPLACEMENTS = {
    "\u2010": "-",
    "\u2011": "-",
    "\u2012": "-",
    "\u2013": "-",
    "\u2014": "-",
    "\u2212": "-",
    "\u00d7": "x",
    "\u2192": "->",
    "\u00a0": " ",
}


def docx_to_pdf(source: bytes | str | Path) -> bytes:
    """Turn a Word resume into a simple text PDF. Some forms accept only PDF."""
    if isinstance(source, (str, Path)):
        document = Document(str(source))
    else:
        document = Document(BytesIO(source))
    lines = [_ascii(paragraph.text) for paragraph in document.paragraphs]
    lines = [line for line in lines if line]
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.multi_cell(0, 6, line)
        pdf.ln(1)
    output = pdf.output()
    if isinstance(output, str):
        return output.encode("latin-1")
    return bytes(output)


def _ascii(text: str) -> str:
    cleaned = "".join(_REPLACEMENTS.get(char, char) for char in text)
    return cleaned.encode("latin-1", errors="replace").decode("latin-1").strip()
