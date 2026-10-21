import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

from job_assistant.nodes import write_cover_letter as letters


class CoverLetterFileTests(unittest.TestCase):
    def test_docx_keeps_email_and_leaves_out_the_phone(self):
        text = letters.finish_letter("I can start immediately.\n+372 5605 7387")
        blob = letters.letter_to_docx(text)
        self.assertTrue(blob.startswith(b"PK"))
        body = "\n".join(p.text for p in Document(BytesIO(blob)).paragraphs)
        self.assertIn("I can start immediately.", body)
        self.assertIn(letters.APPLICANT_EMAIL, body)
        self.assertNotIn("+372", body)
        self.assertNotIn("5605", body)

    def test_persist_cover_writes_a_docx(self):
        original = letters.COVER_LETTERS_PATH
        try:
            with tempfile.TemporaryDirectory() as tmp:
                letters.COVER_LETTERS_PATH = Path(tmp) / "cover_letters.json"
                saved = letters.persist_cover(
                    {
                        "job_id": "job-1",
                        "company": "Perena",
                        "cover_letter": "Hello from the letter.",
                    }
                )
                path = Path(saved["cover_path"])
                self.assertEqual(path.suffix, ".docx")
                self.assertTrue(path.read_bytes().startswith(b"PK"))
                body = "\n".join(p.text for p in Document(path).paragraphs)
                self.assertIn("Hello from the letter.", body)
                self.assertNotIn("+372", body)
        finally:
            letters.COVER_LETTERS_PATH = original


class ResumePhoneTests(unittest.TestCase):
    def test_resume_strip_removes_the_form_phone(self):
        from job_assistant.services.cv_loader import strip_private_contact

        cleaned = strip_private_contact(
            "Email: a@b.com\nPhone: +372 5605 7387\nCall +372 5605 7387.\nReact engineer."
        )
        self.assertNotIn("5605", cleaned)
        self.assertIn("Call", cleaned)
        self.assertIn("a@b.com", cleaned)


if __name__ == "__main__":
    unittest.main()
